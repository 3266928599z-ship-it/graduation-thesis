# -*- coding: utf-8 -*-
"""生成文献综述 Word 版本。"""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re

SRC = "开题/文献综述-基于三维点云配准的免示教焊接关键技术-新版.md"
OUT = "开题/文献综述-基于三维点云配准的免示教焊接关键技术-新版.docx"

doc = Document()
sec = doc.sections[0]
sec.page_height = Cm(29.7)
sec.page_width = Cm(21.0)
sec.top_margin = Cm(2.54)
sec.bottom_margin = Cm(2.54)
sec.left_margin = Cm(2.7)
sec.right_margin = Cm(2.7)

body_font = doc.styles["Normal"].font
body_font.name = "Times New Roman"
body_font.size = Pt(12)
ef = doc.styles["Normal"].element.rPr.rFonts
ef.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ascii", "Times New Roman")
ef.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}hAnsi", "Times New Roman")
ef.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}cs", "Times New Roman")

pstyle = doc.styles["Normal"].paragraph_format
pstyle.space_before = Pt(0)
pstyle.space_after = Pt(0)
pstyle.line_spacing = 1.5
pstyle.alignment = WD_ALIGN_PARAGRAPH.LEFT

def add_cjk_font(run, font_name):
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(r"{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts")
    if rfonts is None:
        rfonts = rpr.makeelement(r"{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts", {})
        rpr.append(rfonts)
    rfonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", font_name)

def add_line(text, bold=False, size=12, align="left"):
    para = doc.add_paragraph()
    para.paragraph_format.line_spacing = 1.5
    para.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
    run = para.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Times New Roman"
    add_cjk_font(run, "黑体" if bold else "仿宋_GB2312")
    return para

def add_heading(text, level):
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    para.paragraph_format.space_after = Pt(6 if level == 1 else 4)
    para.paragraph_format.line_spacing = 1.5
    run = para.add_run(text)
    if level == 1:
        run.font.size = Pt(18)
        add_cjk_font(run, "黑体")
    elif level == 2:
        run.font.size = Pt(16)
        add_cjk_font(run, "黑体")
    else:
        run.font.size = Pt(15)
        add_cjk_font(run, "黑体")
    run.bold = True
    return para

with open(SRC, encoding="utf-8") as f:
    lines = f.read().splitlines()

for line in lines:
    line = line.rstrip()
    if not line.strip():
        continue
    m = re.match(r"^(#{1,4})\s+(.+)$", line)
    if m:
        level = len(m.group(1))
        add_heading(m.group(2), level)
        continue
    # 引用块 > ... 作为斜体或普通段落
    m = re.match(r"^>\s*(.+)$", line)
    if m:
        add_line(m.group(1), bold=False, size=10.5)
        continue
    # 分隔线 ---
    if line.strip() == "---":
        continue
    # 普通段落：粗体 **text** 标记
    bold = False
    if line.startswith("**") and line.endswith("**"):
        bold = True
        line = line[2:-2]
    add_line(line, bold=bold)

doc.save(OUT)
print("saved", OUT)
